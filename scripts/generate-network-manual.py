#!/usr/bin/env python3
"""Generate the editable DOCX and print-ready PDF from the same reviewed source."""
import re
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "generated-artifacts" / "manual-red"
SOURCE = ROOT / "docs" / "manual-funcionamiento-red.md"
LOGO = ROOT / "artifacts" / "web" / "src" / "assets" / "logo.png"
TITLE = "Manual de funcionamiento de la red"
SUBTITLE = "Coordinación de Formación Profesional\nFamilia de Administración y Gestión"
EDITION = "Edición 1.0 · 21 de octubre de 2026"
NAVY = "#12355B"
BLUE = "#155EEF"
PALE = "#EFF5FC"


def parse():
    sections = []
    for part in re.split(r"^# ", SOURCE.read_text(), flags=re.M)[1:]:
        title, body = part.split("\n", 1)
        blocks, lines, i = [], body.strip().splitlines(), 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue
            if line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r"[:\-\s]+", c) for c in row):
                        rows.append(row)
                    i += 1
                blocks.append(("table", rows))
                continue
            if line.startswith("## "):
                blocks.append(("heading", line[3:]))
            elif line.startswith("> "):
                blocks.append(("note", line[2:]))
            elif line.startswith("- "):
                blocks.append(("bullet", line[2:]))
            elif re.match(r"^\d+\.\s", line):
                blocks.append(("step", line))
            else:
                blocks.append(("text", line))
            i += 1
        sections.append((title.strip(), blocks))
    assert len(sections) == 18, "Review the contents page when adding sections"
    return sections


def inline(text):
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escape(text))


def rich_docx(paragraph, text):
    for index, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            paragraph.add_run(chunk).bold = index % 2 == 1


def cell_shade(cell, color):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(element)


def make_docx(sections):
    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin, page.bottom_margin = Cm(1.9), Cm(1.8)
    page.left_margin = page.right_margin = Cm(2)
    page.different_first_page_header_footer = True
    for name in ("Normal", "Body Text", "List Bullet"):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 29), ("Heading 1", 20), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string(NAVY[1:])
        style.paragraph_format.space_after = Pt(8)
    header = page.header.paragraphs[0]
    header.text = "RED DE COORDINACIÓN DE FP  |  MANUAL OPERATIVO"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(NAVY[1:])
    footer = page.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Edición 1.0  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)

    doc.add_paragraph("\n")
    doc.add_picture(str(LOGO), width=Cm(12))
    doc.add_paragraph("\n")
    doc.add_paragraph(TITLE, "Title")
    doc.add_paragraph(SUBTITLE, "Subtitle")
    doc.add_paragraph("Guía de organización, participación y trabajo compartido")
    doc.add_paragraph(EDITION)
    doc.add_paragraph("Basado en las funciones disponibles en Coordina ADG.")
    doc.add_paragraph("Referencia institucional: coordinacionfp.org")
    doc.add_page_break()
    doc.add_heading("Cómo utilizar este manual", 1)
    doc.add_paragraph("Lectura inicial: capítulos 1 a 5. Uso cotidiano: capítulos 6 a 13. Seguimiento y seguridad: capítulos 14 a 16. Aplicación práctica: listas y modelos de los capítulos 17 y 18.")
    doc.add_paragraph("Se distinguen las reglas implementadas en la aplicación de las pautas organizativas que debe adoptar la red. Los permisos efectivos y las políticas de la entidad prevalecen.")
    doc.add_heading("Contenido", 2)
    for title, _ in sections:
        doc.add_paragraph(title)
    doc.add_paragraph("El documento editable utiliza estilos de título para facilitar la navegación y futuras revisiones.")

    for title, blocks in sections:
        doc.add_page_break()
        doc.add_heading(title, 1)
        for kind, value in blocks:
            if kind == "heading":
                doc.add_heading(value, 2)
            elif kind == "table":
                table = doc.add_table(rows=0, cols=len(value[0]))
                table.style = "Light Shading Accent 1"
                for r, row in enumerate(value):
                    cells = table.add_row().cells
                    for cell, text in zip(cells, row):
                        rich_docx(cell.paragraphs[0], text)
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(9)
                            if r == 0:
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                        if r == 0:
                            cell_shade(cell, NAVY[1:])
                    props = table.rows[-1]._tr.get_or_add_trPr()
                    avoid_split = OxmlElement("w:cantSplit")
                    props.append(avoid_split)
                    if r == 0:
                        repeat = OxmlElement("w:tblHeader")
                        props.append(repeat)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            else:
                p = doc.add_paragraph(style="List Bullet" if kind == "bullet" else "Normal")
                rich_docx(p, value)
                if kind == "note":
                    p.paragraph_format.left_indent = Cm(0.25)
                    for run in p.runs:
                        run.italic = True
                        run.font.color.rgb = RGBColor.from_string(NAVY[1:])
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Organización y procedimientos de la red de coordinación de FP"
    doc.core_properties.author = "Red de coordinación de Formación Profesional"
    doc.core_properties.version = "1.0"
    doc.save(OUT / "Manual_funcionamiento_red_FP.docx")


class ManualPDF(SimpleDocTemplate):
    def afterFlowable(self, flowable):
        if hasattr(flowable, "bookmark"):
            self.canv.bookmarkPage(flowable.bookmark)
            self.canv.addOutlineEntry(flowable.getPlainText(), flowable.bookmark, level=0)


def make_pdf(sections):
    fonts = Path("/usr/share/fonts/truetype/dejavu")
    pdfmetrics.registerFont(TTFont("Manual", str(fonts / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("Manual-Bold", str(fonts / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFontFamily("Manual", normal="Manual", bold="Manual-Bold", italic="Manual", boldItalic="Manual-Bold")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("BodyManual", fontName="Manual", fontSize=9.3, leading=14, spaceAfter=7, textColor=colors.HexColor("#263747")))
    styles.add(ParagraphStyle("ChapterManual", parent=styles["BodyManual"], fontName="Manual-Bold", fontSize=19, leading=25, textColor=colors.HexColor(NAVY), spaceAfter=18))
    styles.add(ParagraphStyle("SubManual", parent=styles["BodyManual"], fontName="Manual-Bold", fontSize=11.5, leading=16, textColor=colors.HexColor(NAVY), spaceBefore=7, spaceAfter=6, keepWithNext=True))
    styles.add(ParagraphStyle("CellManual", parent=styles["BodyManual"], fontSize=8.3, leading=12, spaceAfter=0))
    styles.add(ParagraphStyle("HeadCellManual", parent=styles["CellManual"], fontName="Manual-Bold", textColor=colors.white))
    styles.add(ParagraphStyle("NoteManual", parent=styles["BodyManual"], fontSize=8.5, leading=12.5, textColor=colors.HexColor(NAVY), borderColor=colors.HexColor("#C4D7EF"), borderWidth=0.6, borderPadding=9, backColor=colors.HexColor(PALE), spaceBefore=7, spaceAfter=9))
    styles.add(ParagraphStyle("BulletManual", parent=styles["BodyManual"], leftIndent=11, firstLineIndent=-10, spaceAfter=5))
    styles.add(ParagraphStyle("CoverManual", parent=styles["ChapterManual"], fontSize=31, leading=39, spaceAfter=22))
    styles.add(ParagraphStyle("CoverSubManual", parent=styles["BodyManual"], fontSize=15, leading=23, spaceAfter=20, textColor=colors.HexColor(NAVY)))
    doc = ManualPDF(str(OUT / "Manual_funcionamiento_red_FP.pdf"), pagesize=A4,
                    rightMargin=52, leftMargin=52, topMargin=57, bottomMargin=48,
                    title=TITLE, author="Red de coordinación de Formación Profesional")
    flow = [Spacer(1, 40)]
    with PILImage.open(LOGO) as image:
        width, height = image.size
    flow += [Image(str(LOGO), width=345, height=345 * height / width, hAlign="LEFT"),
             Spacer(1, 48), Paragraph(TITLE, styles["CoverManual"]),
             Paragraph(SUBTITLE.replace("\n", "<br/>"), styles["CoverSubManual"]),
             Paragraph("Guía de organización, participación<br/>y trabajo compartido", styles["BodyManual"]),
             Spacer(1, 35), Paragraph(EDITION, styles["BodyManual"]),
             Paragraph("Basado en las funciones disponibles en Coordina ADG.", styles["BodyManual"]),
             Paragraph("Referencia institucional: <link href='https://coordinacionfp.org' color='#155EEF'>coordinacionfp.org</link>", styles["BodyManual"]),
             PageBreak(), Paragraph("Cómo utilizar este manual", styles["ChapterManual"])]
    flow += [Paragraph("Lectura inicial: capítulos 1 a 5. Uso cotidiano: capítulos 6 a 13. Seguimiento y seguridad: capítulos 14 a 16. Aplicación práctica: listas y modelos de los capítulos 17 y 18.", styles["BodyManual"]),
             Paragraph("Se distinguen las reglas implementadas en la aplicación de las pautas organizativas que debe adoptar la red. Los permisos efectivos y las políticas de la entidad prevalecen.", styles["NoteManual"]),
             Spacer(1, 8), Paragraph("Contenido", styles["SubManual"])]
    for index, (title, _) in enumerate(sections, 1):
        flow.append(Paragraph(f'<link href="#section-{index:02d}" color="{NAVY}">{escape(title)}</link>', styles["BodyManual"]))
    flow.append(Spacer(1, 12))
    flow.append(Paragraph("El índice del PDF es navegable. El documento editable utiliza estilos de título para facilitar las revisiones.", styles["BodyManual"]))

    for index, (title, blocks) in enumerate(sections, 1):
        heading = Paragraph(escape(title), styles["ChapterManual"])
        heading.bookmark = f"section-{index:02d}"
        flow += [PageBreak(), heading]
        for kind, value in blocks:
            if kind == "table":
                count = len(value[0])
                widths = [doc.width * f for f in ((0.34, 0.66) if count == 2 else (0.23, 0.27, 0.50))]
                rows = [[Paragraph(inline(text), styles["HeadCellManual"] if r == 0 else styles["CellManual"])
                         for text in row] for r, row in enumerate(value)]
                table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FA")]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ]))
                flow += [table, Spacer(1, 10)]
            elif kind == "heading":
                flow.append(Paragraph(inline(value), styles["SubManual"]))
            elif kind == "note":
                flow.append(Paragraph(inline(value), styles["NoteManual"]))
            elif kind == "bullet":
                flow.append(Paragraph("• " + inline(value), styles["BulletManual"]))
            else:
                flow.append(Paragraph(inline(value), styles["BodyManual"]))

    def decorate(canvas, document):
        canvas.saveState()
        if document.page > 1:
            canvas.setStrokeColor(colors.HexColor("#D7E2F0"))
            canvas.line(52, A4[1] - 39, A4[0] - 52, A4[1] - 39)
            canvas.setFont("Manual-Bold", 7)
            canvas.setFillColor(colors.HexColor(NAVY))
            canvas.drawString(52, A4[1] - 31, "RED DE COORDINACIÓN DE FP  |  MANUAL OPERATIVO")
        canvas.setFont("Manual", 7.5)
        canvas.setFillColor(colors.HexColor("#596A7C"))
        canvas.drawString(52, 29, "Edición 1.0 · Octubre de 2026")
        canvas.drawRightString(A4[0] - 52, 29, str(document.page))
        canvas.restoreState()
    doc.build(flow, onFirstPage=decorate, onLaterPages=decorate)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    sections = parse()
    make_docx(sections)
    make_pdf(sections)
    print("Generados DOCX y PDF en generated-artifacts/manual-red/")
